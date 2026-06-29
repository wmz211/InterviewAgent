"""Structural and cross-report audits for generated InterviewAgent DOCX files."""

import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from scripts.generate_project_management_reports import REPORTS
from scripts.project_reports_data import PROJECT, money


REQUIRED = {
    "01_项目建议书.docx": ["执行摘要", "问题与机会", "拟建方案", "立项结论"],
    "02_可行性分析报告.docx": ["技术可行性", "经济可行性", "敏感性分析", "结论与约束条件"],
    "03_项目计划书.docx": ["工作分解结构（WBS）", "资源计划", "进度计划", "成本计划", "质量计划", "风险计划"],
    "04_项目估算报告.docx": ["资源估算", "工期估算", "成本估算", "PERT 三点估算"],
    "05_项目进度和成本控制报告.docx": ["20 周总进度安排", "成本基线与阶段预算", "挣值管理规则", "偏差阈值与纠偏流程"],
    "06_项目质量管理报告.docx": ["质量保证活动", "质量控制与测试策略", "AI 与 RAG 专项质量", "缺陷分级与根因分析机制"],
    "07_项目风险管理报告.docx": ["风险登记册", "触发器与应急计划", "风险复审与问题转化安排"],
    "08_项目团队建设和管理报告.docx": ["团队目标与组织结构", "技能矩阵与知识备份", "团队工作协议", "冲突与决策管理"],
    "09_项目总结报告.docx": ["项目目标与范围总结", "管理基线总结", "项目管理重点", "总结结论"],
}


def extract_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def audit_table_geometry(doc, filename):
    errors = []
    for index, table in enumerate(doc.tables, 1):
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        if tbl_w is None or tbl_w.get(qn("w:w")) != "9360" or tbl_w.get(qn("w:type")) != "dxa":
            errors.append(f"{filename}: table {index} width is not 9360 DXA")
        if tbl_ind is None or tbl_ind.get(qn("w:w")) != "120":
            errors.append(f"{filename}: table {index} indent is not 120 DXA")
        grid_widths = [int(c.get(qn("w:w"))) for c in table._tbl.tblGrid]
        if sum(grid_widths) != 9360:
            errors.append(f"{filename}: table {index} grid totals {sum(grid_widths)}")
        for row in table.rows:
            if row._tr.get_or_add_trPr().find(qn("w:trHeight")) is not None:
                errors.append(f"{filename}: table {index} uses fixed row height")
            widths = []
            for cell in row.cells:
                tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
                widths.append(int(tc_w.get(qn("w:w"))) if tc_w is not None else 0)
            if widths != grid_widths:
                errors.append(f"{filename}: table {index} cell widths do not match grid")
                break
    return errors


def audit_file(path):
    filename = path.name
    errors = []
    doc = Document(path)
    text = extract_text(doc)
    for required in [PROJECT["name"], "作业编制小组", *PROJECT["authors"], "陈昊", "林泽宇", "20 周"]:
        if required not in text:
            errors.append(f"{filename}: missing shared fact {required}")
    for required in REQUIRED.get(filename, []):
        if required not in text:
            errors.append(f"{filename}: missing required content {required}")
    for marker in ("TBD", "TODO", "待填写", "请补充", "XX"):
        if marker in text:
            errors.append(f"{filename}: placeholder marker {marker}")
    for forbidden in (
        "最终模拟实际",
        "第 14 周状态日",
        "结项状态",
        "实际成本",
        "较基线延期",
        "三人全职",
        "三人团队",
        "王明哲（项目经理",
        "江悦铭（后端",
        "周士斌（前端",
        "21 周",
        "2,400 小时",
        "¥480,810",
    ):
        if forbidden in text:
            errors.append(f"{filename}: forbidden execution/role wording {forbidden}")
    if len(text) < 1500:
        errors.append(f"{filename}: content is unexpectedly short ({len(text)} chars)")
    errors.extend(audit_table_geometry(doc, filename))
    with zipfile.ZipFile(path) as package:
        names = set(package.namelist())
        if "word/header1.xml" not in names or "word/footer1.xml" not in names:
            errors.append(f"{filename}: missing header/footer")
        footer = package.read("word/footer1.xml").decode("utf-8")
        if " PAGE " not in footer:
            errors.append(f"{filename}: missing PAGE field")
    return errors


def audit_directory(directory):
    directory = Path(directory)
    errors = []
    expected = [name for name, _ in REPORTS]
    found = sorted(path.name for path in directory.glob("*.docx"))
    if found != sorted(expected):
        errors.append(f"file set mismatch: expected {sorted(expected)}, found {found}")
    for filename in expected:
        path = directory / filename
        if path.exists():
            errors.extend(audit_file(path))
    return errors


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("output/project_management_reports")
    result = audit_directory(target)
    if result:
        print("AUDIT FAILED")
        for item in result:
            print(f"- {item}")
        raise SystemExit(1)
    print(f"AUDIT OK: 9 reports; BAC=RMB {PROJECT['bac']:,}; planned=20 weeks; fictional team=13")
