from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "InterviewAgent实验报告-王明哲.docx"


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
BORDER = "D9E2EC"
TEXT = "1F2933"
MUTED = "5C6670"
ACCENT = "C96214"


def set_run_font(run, size: float | None = None, bold: bool = False, color: str | None = None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")


def set_cell_text(cell, text: str, bold: bool = False, color: str = TEXT, size: float = 10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def add_para(doc, text: str = "", style: str | None = None, bold_prefix: str | None = None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, 11, True, TEXT)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, 11, False, TEXT)
    else:
        run = p.add_run(text)
        set_run_font(run, 11, False, TEXT)
    return p


def add_bullet(doc, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run, 11, False, TEXT)
    return p


def add_number(doc, text: str):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run, 11, False, TEXT)
    return p


def add_heading(doc, text: str, level: int = 1):
    p = doc.add_heading(level=level)
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12 if level == 2 else 8)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 6 if level == 2 else 4)
    run = p.add_run(text)
    set_run_font(run, 16 if level == 1 else 13 if level == 2 else 12, True, BLUE if level < 3 else DARK_BLUE)
    return p


def add_callout(doc, title: str, body: str):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, 10.5, True, DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.1
    r2 = p2.add_run(body)
    set_run_font(r2, 10.5, False, TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_figure_placeholder(doc, number: int, title: str, instruction: str):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "FFF7ED")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"【截图占位：图 {number}  {title}】")
    set_run_font(r, 11, True, ACCENT)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(instruction)
    set_run_font(r2, 10, False, MUTED)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    rr = cap.add_run(f"图 {number}  {title}")
    set_run_font(rr, 10, False, MUTED)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, widths)
    for i, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], LIGHT_FILL)
        set_cell_text(table.rows[0].cells[i], header, bold=True, color=DARK_BLUE, size=10)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value), size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def configure_doc(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def add_footer(doc: Document):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("InterviewAgent 实验报告")
    set_run_font(run, 9, False, MUTED)


def build():
    doc = Document()
    configure_doc(doc)
    add_footer(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(72)
    title.paragraph_format.space_after = Pt(8)
    r = title.add_run("InterviewAgent 智能面试系统实验报告")
    set_run_font(r, 22, True, DARK_BLUE)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(32)
    r = subtitle.add_run("基于 LangGraph、GraphRAG、实时语音与结构化评测的 AI 面试 Agent 系统")
    set_run_font(r, 12, False, MUTED)

    meta_rows = [
        ("实验题目", "InterviewAgent 智能面试系统设计与实现"),
        ("学生姓名", "王明哲"),
        ("学号", "2327406062"),
        ("项目目录", r"D:\Homework\InterviewAgent"),
        ("报告版本", "实验报告初稿，可继续插入截图与教师要求的封面信息"),
    ]
    add_table(doc, ["项目", "内容"], meta_rows, [1.45, 5.05])

    add_callout(
        doc,
        "摘要",
        "本实验围绕 AI 面试 Agent 系统的设计、实现与验证展开。系统以 FastAPI 作为后端服务入口，"
        "使用 LangGraph 将面试官行为拆分为开场、简历深挖、JD 技术考察、编程题、HR 行为面与总结评估等阶段，"
        "并结合知识图谱、BM25、向量检索、RRF 融合排序、实时语音和 JWT 多用户认证，形成可交互、可追踪、可评估的模拟面试平台。",
    )

    doc.add_section(WD_SECTION.NEW_PAGE)

    add_heading(doc, "一、实验题目", 1)
    add_para(doc, "InterviewAgent 智能面试系统设计与实现。")
    add_para(
        doc,
        "本实验以“AI 虚拟面试官”为目标对象，完成一个支持简历与岗位描述输入、自动生成面试追问、"
        "多阶段推进面试流程、实时文本/语音交互、结构化评估和管理员后台统计的综合项目。"
    )

    add_heading(doc, "二、实验目的", 1)
    purposes = [
        "掌握基于 FastAPI 的 Web 后端接口设计、路由组织、静态页面托管、CORS 配置和运行时安全校验方法。",
        "理解并实践 LangGraph 在多阶段 Agent 工作流中的应用，将复杂面试流程拆分为可编排、可调试、可观测的节点。",
        "掌握 RAG 与 GraphRAG 的工程实现思路，能够结合知识图谱、BM25、向量检索与 RRF 排序生成更有层次的技术追问。",
        "实践多用户认证、JWT 鉴权、SQLite 持久化、Redis/内存会话状态管理和历史报告查询等系统能力。",
        "掌握实时交互系统的基本实现方式，包括 SSE 流式输出、WebSocket 音频链路、ASR 语音识别和 TTS 播放。",
        "通过测试、快速验证脚本和管理员后台，提升项目的可维护性、可展示性和上线前可检查能力。",
    ]
    for item in purposes:
        add_bullet(doc, item)

    add_heading(doc, "三、实验任务（内容）", 1)
    add_para(
        doc,
        "本实验任务不是实现单一聊天机器人，而是围绕“可编排的 AI 面试官”构建完整系统。"
        "系统需要从候选人简历和岗位 JD 出发，自动组织面试流程，并在面试结束后输出可解释的能力评估。"
    )
    task_rows = [
        ("用户与权限", "实现注册、登录、JWT 鉴权、用户会话隔离，以及管理员后台访问控制。"),
        ("资料上传", "支持上传简历文件、输入 JD、复用历史简历/JD，并创建面试 session。"),
        ("面试流程", "支持技术面与 HR 面两条路径，按阶段推进，也支持指定阶段专项练习。"),
        ("Agent 编排", "使用 LangGraph 将面试节点拆分为 greeting、resume_dive、jd_tech、coding_test、hr_*、wrap_up。"),
        ("RAG/GraphRAG", "构建技术知识图谱、题库和向量库，融合检索结果生成技术追问链。"),
        ("实时交互", "支持文本流式输出、WebSocket 语音输入、ASR 转写和 TTS 播放。"),
        ("结果评估", "记录每轮问答，按阶段生成评分、薄弱项、参考答案和学习建议。"),
        ("后台展示", "管理员登录后查看覆盖领域、题库数量、知识节点、用户数、会话数和运行分布。"),
        ("验证测试", "补充单元测试、静态前端回归测试和快速验证脚本，降低回归风险。"),
    ]
    add_table(doc, ["任务模块", "具体内容"], task_rows, [1.55, 4.95])

    add_heading(doc, "四、实验过程", 1)
    add_heading(doc, "4.1 需求分析与总体方案", 2)
    add_para(
        doc,
        "项目的核心需求可以概括为：让系统像真实面试官一样，先理解候选人的简历和岗位要求，再按面试阶段逐步提问，"
        "根据候选人的回答继续追问，最终形成评价报告。相比普通问答系统，本项目更加关注流程控制、上下文隔离、"
        "追问质量、结果评估和多用户安全边界。"
    )
    add_para(
        doc,
        "总体方案采用“静态前端 + FastAPI API 层 + JWT 鉴权 + 会话状态存储 + LangGraph 工作流 + RAG/GraphRAG 检索增强”的架构。"
        "前端负责登录、上传、面试对话、语音控制和历史结果展示；后端负责用户认证、文件处理、面试状态推进、检索增强、"
        "消息持久化和评估输出；管理员后台用于展示系统覆盖范围和运行数据。"
    )
    add_figure_placeholder(
        doc,
        1,
        "系统首页/登录界面",
        "建议截图位置：http://127.0.0.1:8000 或 8001 的登录页面，展示 InterviewAgent 的整体入口。",
    )

    add_heading(doc, "4.2 开发环境与项目结构", 2)
    env_rows = [
        ("后端框架", "FastAPI、Pydantic、SQLAlchemy async、Uvicorn"),
        ("Agent 与 LLM", "LangGraph、LangChain、Qwen/DashScope OpenAI-compatible API"),
        ("检索与知识库", "Neo4j/NetworkX、ChromaDB、BM25、SentenceTransformer、RRF"),
        ("语音链路", "WebSocket、DashScope ASR、DashScope TTS"),
        ("存储", "SQLite、Redis 可选、JSON session archive"),
        ("测试", "pytest、unittest、compileall、scripts/verify_fast.py"),
    ]
    add_table(doc, ["类别", "使用技术"], env_rows, [1.6, 4.9])
    add_para(
        doc,
        "项目主要目录包括 app/api、app/auth、app/core、app/db、app/agents、app/rag、static、data、scripts 和 tests。"
        "其中 app/api 负责接口路由，app/core 负责面试状态与工作流，app/rag 负责知识图谱和向量检索，static 提供前端界面，"
        "tests 中包含阶段流转、会话安全、RAG 评估、前端回归等测试。"
    )

    add_heading(doc, "4.3 系统架构设计", 2)
    add_para(
        doc,
        "系统入口为 app/main.py。启动时加载 .env 配置，完成运行时安全校验，注册 API 路由并托管 static 目录。"
        "API 聚合路由位于 app/api/router.py，分别挂载 auth、upload、interview、ws 和 admin 模块。"
        "用户请求进入后首先经过 JWT 鉴权，再根据功能进入上传、面试、语音或后台统计链路。"
    )
    add_para(
        doc,
        "面试状态由 session_id 标识。上传简历和 JD 后，后端创建初始状态，写入数据库 session 元数据，并将活跃状态放入会话存储。"
        "每次候选人回答后，后端把回答追加到 messages 中，调用 LangGraph 推进当前节点，产生新的面试官问题或阶段切换结果，"
        "同时将对话 transcript 持久化，便于后续生成报告和历史查询。"
    )
    add_figure_placeholder(
        doc,
        2,
        "系统架构图或 README 架构图",
        "建议截图 README 中的 Mermaid 架构图，或在浏览器/API 文档中展示模块关系。",
    )

    add_heading(doc, "4.4 用户认证与会话安全实现", 2)
    add_para(
        doc,
        "认证模块位于 app/api/endpoints/auth.py 与 app/auth。用户注册时系统检查 email 和 username 是否重复，"
        "使用 passlib 对密码进行哈希后写入 users 表；登录时校验密码并生成 JWT。受保护接口通过 get_current_user 解析 Bearer token，"
        "确保只有有效且激活的用户能够访问自己的面试资源。"
    )
    add_para(
        doc,
        "会话安全的重点是防止不同用户之间越权读取。项目通过 user_id 绑定 interview_sessions，并在读取报告、修改阶段、访问历史数据时校验 session 所属用户。"
        "测试中覆盖了 user_id 匹配、不同用户拒绝、历史遗留 session 无 owner 时拒绝等边界情况。"
    )

    add_heading(doc, "4.5 简历/JD 上传与材料复用", 2)
    add_para(
        doc,
        "上传模块支持候选人上传 PDF/DOCX 简历并输入岗位 JD。系统从简历中抽取文本，结合 JD 创建面试 session。"
        "为了提高重复练习效率，前端提供历史简历和历史 JD 下拉选择，后端提供 reuse-options 接口，使用户可以复用之前的材料，"
        "避免每次练习都重新上传同一份文件。"
    )
    add_figure_placeholder(
        doc,
        3,
        "简历上传、JD 输入与面试模式选择界面",
        "建议截图上传页，包含技术面/HR 面选择、历史简历/JD 复用、简历上传区和 JD 文本框。",
    )

    add_heading(doc, "4.6 LangGraph 多阶段面试工作流", 2)
    add_para(
        doc,
        "本项目将面试官拆分为多个节点，而不是让一个大 prompt 处理全部逻辑。技术面流程包括 greeting、resume_dive、jd_tech、coding_test 和 wrap_up；"
        "HR 面流程包括 hr_self_intro、hr_behavioral、hr_career 和 wrap_up。每个节点只负责本阶段的提问、追问和局部评分，"
        "通过状态字段 current_node、phase_turn_count、should_transition、interview_complete 等控制阶段流转。"
    )
    add_para(
        doc,
        "这种拆分方式有三个优点：第一，面试流程可控，避免模型无限发散；第二，不同阶段可以使用不同的 prompt、工具和评分规则；"
        "第三，便于测试和定位问题，例如单独验证 resume_dive 是否围绕简历项目追问，jd_tech 是否结合岗位技术点提问。"
    )
    phase_rows = [
        ("greeting", "开场确认", "确认候选人身份、面试模式和自我介绍。"),
        ("resume_dive", "简历深挖", "围绕项目经历、职责边界、技术难点和结果指标追问。"),
        ("jd_tech", "JD 技术考察", "抽取岗位技术点，结合 GraphRAG 形成层次化追问。"),
        ("coding_test", "算法编程", "考察算法思路、复杂度、边界条件和代码表达。"),
        ("hr_*", "HR 轨道", "覆盖自我介绍、行为面试、职业规划和动机稳定性。"),
        ("wrap_up", "总结评估", "输出阶段评分、薄弱项、参考答案和学习建议。"),
    ]
    add_table(doc, ["节点", "阶段", "作用"], phase_rows, [1.3, 1.45, 3.75])

    add_heading(doc, "4.7 GraphRAG 与题库检索实现", 2)
    add_para(
        doc,
        "GraphRAG 是本项目技术追问质量的核心。系统维护 data/tech_knowledge_graph.json，包含 Domain、Tech、Concept、Component 和 Pitfall 等节点类型，"
        "并通过 BELONGS_TO、HAS_COMPONENT、LEADS_TO、HAS_PITFALL、REQUIRES、RELATED_TO 等关系描述知识之间的依赖、组成和追问路径。"
        "Neo4j 可用时使用图数据库查询；本地快速开发时可回退到 NetworkX。"
    )
    add_para(
        doc,
        "在 JD 技术考察阶段，系统先抽取岗位文本中的技术实体，再进行 BM25 精确词召回和向量语义召回，最后用 RRF 融合多个排名。"
        "命中技术节点后，系统沿 LEADS_TO、HAS_COMPONENT、HAS_PITFALL 等边扩展追问，使问题从概念解释自然深入到原理、工程实践、风险和评估指标。"
    )
    add_para(
        doc,
        "题库位于 data/question_bank，包括大模型/RAG/Agent 问题和算法编程问题。题目带有 topic、difficulty 和 question_type 等字段，"
        "用于在不同阶段提供可复用问题和参考答案。"
    )

    add_heading(doc, "4.8 前端交互与实时语音链路", 2)
    add_para(
        doc,
        "前端页面位于 static/index.html，采用单页多视图方式组织认证、上传、面试、报告和历史记录。登录成功后 token 保存到 localStorage，"
        "后续请求自动携带 Authorization 头。面试页面提供阶段导航、候选人输入框、流式回答展示、语音按钮和 TTS 播放控制。"
    )
    add_para(
        doc,
        "语音交互通过 WebSocket 实现。前端把音频流发送到 /api/v1/ws/audio/{session_id}?token={jwt_token}，后端校验 token 后进行 ASR 转写，"
        "最终把识别文本追加到输入框。TTS 用于播放面试官问题，使系统更接近真实面试场景。"
    )
    add_figure_placeholder(
        doc,
        4,
        "面试对话与阶段导航界面",
        "建议截图一次完整技术面或 HR 面对话，展示左侧阶段、候选人回答和面试官追问。",
    )

    add_heading(doc, "4.9 管理员后台实现", 2)
    add_para(
        doc,
        "为便于展示系统覆盖范围，本实验新增管理员后台页面 static/admin.html，并在后端新增 /api/v1/admin/overview 接口。"
        "后台不单独维护一套账号体系，而是复用现有用户登录，再通过 ADMIN_EMAILS 白名单判断是否具有管理员权限。"
        "管理员可以查看用户数、面试会话数、对话消息数、覆盖领域数、知识图谱节点与关系数量、题库统计、面试阶段和会话分布。"
    )
    add_para(
        doc,
        "后台统计数据来源包括 SQLite 数据库、data/tech_knowledge_graph.json 和 data/question_bank/*.json。"
        "这种设计避免引入复杂迁移，也使后台能随着知识图谱和题库文件变化自动更新展示内容。"
    )
    add_figure_placeholder(
        doc,
        5,
        "管理员后台系统概览页面",
        "建议截图 http://127.0.0.1:8001/admin 登录后的概览页，重点展示指标卡片和覆盖领域。",
    )

    add_heading(doc, "4.10 测试与验证过程", 2)
    add_para(
        doc,
        "实验过程中使用 compileall、unittest、pytest 和 scripts/verify_fast.py 进行验证。快速验证脚本会检查 app、scripts、tests 的语法，"
        "运行关键单元测试，并扫描 README、.env.example、Compose 和 docs 中是否存在明显敏感值。"
    )
    test_rows = [
        ("语法检查", "python -m compileall -q app scripts tests", "通过"),
        ("后台测试", "pytest tests/test_admin_dashboard.py -q", "6 passed"),
        ("前端回归", "pytest tests/test_static_frontend_regressions.py -q", "通过"),
        ("快速验证", "python scripts/verify_fast.py", "Fast verification passed"),
        ("接口验证", "GET /api/v1/admin/overview", "管理员邮箱通过，返回 metrics=6"),
    ]
    add_table(doc, ["验证项", "命令或接口", "结果"], test_rows, [1.45, 3.7, 1.35])
    add_figure_placeholder(
        doc,
        6,
        "测试通过或 API 文档页面",
        "建议截图 pytest/verify_fast 终端输出，或截图 FastAPI /docs 中 admin、auth、interview 接口列表。",
    )

    add_heading(doc, "五、实验结果（展示）", 1)
    add_heading(doc, "5.1 功能实现结果", 2)
    add_para(
        doc,
        "实验最终完成了一个可运行的 AI 面试 Agent 系统。用户可以注册登录，上传简历和 JD，选择技术面或 HR 面，"
        "进入多阶段模拟面试；系统能够根据候选人回答继续追问，并在结束时生成结构化评估结果。"
    )
    result_rows = [
        ("注册登录", "用户可注册、登录并持有 JWT，后端接口按 token 鉴权。"),
        ("面试创建", "上传简历和 JD 后生成 session，保存用户归属、模式和阶段信息。"),
        ("阶段推进", "技术面和 HR 面均可按工作流推进，并支持阶段练习。"),
        ("GraphRAG 追问", "JD 技术阶段可结合知识图谱和混合检索生成层次化问题。"),
        ("实时交互", "文本流式输出、WebSocket 音频、ASR/TTS 链路具备基础可用性。"),
        ("历史与报告", "可查询历史会话和结构化报告，便于复盘。"),
        ("管理员后台", "管理员可查看系统覆盖领域、题库、用户、会话和消息统计。"),
    ]
    add_table(doc, ["功能", "实现结果"], result_rows, [1.45, 5.05])

    add_heading(doc, "5.2 界面展示结果", 2)
    add_para(
        doc,
        "界面展示建议至少插入 6 张截图，以证明系统从登录、材料上传、面试交互、后台统计到测试验证均已完成。"
        "下面的截图占位已经写入正文，后续只需要在 Word 中替换对应占位即可。"
    )
    screenshot_rows = [
        ("图 1", "系统首页/登录界面", "证明用户入口和认证功能。"),
        ("图 2", "系统架构图或 README 架构图", "证明总体设计和模块关系。"),
        ("图 3", "上传简历、JD 与模式选择界面", "证明实验输入与任务配置功能。"),
        ("图 4", "面试对话与阶段导航界面", "证明多阶段面试流程和交互效果。"),
        ("图 5", "管理员后台系统概览页面", "证明覆盖领域、题库和运行统计展示。"),
        ("图 6", "测试通过或 API 文档页面", "证明验证过程和接口完整性。"),
    ]
    add_table(doc, ["图号", "截图内容", "展示目的"], screenshot_rows, [0.8, 2.4, 3.3])

    add_heading(doc, "5.3 性能与可维护性结果", 2)
    add_para(
        doc,
        "项目在本地开发环境下支持轻量运行。Neo4j 不可用时，系统可通过 NetworkX 读取 JSON 图谱完成基础 GraphRAG；Redis 未配置时，"
        "系统使用内存会话存储并在面试完成后持久化 JSON；SQLite 用于保存用户、会话和 transcript。这样的降级策略降低了部署门槛，"
        "适合课程实验展示和本地调试。"
    )
    add_para(
        doc,
        "代码结构上，API、认证、数据库、核心工作流、RAG、音频和前端静态资源分层清晰。新增管理员后台时只需要新增 admin endpoint、"
        "扩展路由、增加静态页面和少量测试，没有破坏现有面试流程，说明系统具备一定可扩展性。"
    )

    add_heading(doc, "六、实验总结", 1)
    add_heading(doc, "6.1 已完成的主要工作", 2)
    for item in [
        "完成了 InterviewAgent 的核心业务闭环：上传材料、创建面试、阶段追问、记录对话、生成评估、查看历史。",
        "完成了技术面和 HR 面两类流程建模，使面试行为从普通问答升级为可控的多阶段 Agent 工作流。",
        "完成了 GraphRAG 技术追问能力，将岗位技术点、知识图谱、题库和混合检索结合起来，提高问题层次感。",
        "完成了多用户认证和会话隔离，使系统具备基本生产安全边界。",
        "完成了管理员后台，能够从管理视角展示系统覆盖领域、知识节点、题库和运行统计。",
        "补充了面向核心逻辑的测试和快速验证脚本，使项目更适合持续迭代。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "6.2 实验中遇到的问题与解决方法", 2)
    issue_rows = [
        ("LLM 输出不可控", "通过阶段节点、最大轮次、状态字段和 fallback 控制流程，避免无边界对话。"),
        ("技术追问容易停留在表层", "引入知识图谱关系和 RRF 融合检索，让问题沿概念、组件、风险和工程实践深入。"),
        ("会话数据容易丢失", "将 session 元数据和 transcript 写入 SQLite，活跃状态使用内存/Redis，完成后持久化 JSON。"),
        ("多用户越权风险", "所有 session 相关接口校验 user_id，测试覆盖 owner 匹配和拒绝场景。"),
        ("本地依赖较复杂", "提供 NetworkX 和内存存储降级路径，减少 Neo4j/Redis 未启动时的阻塞。"),
        ("管理员后台权限更新不生效", "将管理员邮箱白名单从 .env 中读取，并在调试中确认新服务加载配置。"),
    ]
    add_table(doc, ["问题", "解决方法"], issue_rows, [2.0, 4.5])

    add_heading(doc, "6.3 不足与改进方向", 2)
    for item in [
        "进一步完善报告持久化，将最终评估结果也完整写入数据库，而不是只依赖活跃状态或 JSON 归档。",
        "增加更系统的 RAG 评估集，量化 Recall@K、MRR、追问相关性和 GraphRAG 相比普通 RAG 的提升。",
        "补充端到端浏览器自动化测试，覆盖登录、上传、面试、后台和历史记录等完整用户路径。",
        "优化 ASR/TTS 的异常处理、重试、超时和用户提示，提升实时语音链路稳定性。",
        "管理员后台后续可增加用户管理、会话详情、题库管理、知识图谱编辑和评估质量监控。",
        "生产环境还应补充限流、上传大小限制、审计日志、HTTPS、密钥轮换和更严格的 CORS 配置。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "6.4 个人收获", 2)
    add_para(
        doc,
        "通过本次实验，我对 AI Agent 系统的工程化实现有了更完整的认识。一个可用的 Agent 项目不仅需要大模型能力，"
        "还需要流程编排、状态管理、检索增强、权限控制、数据持久化、前端交互、测试验证和运维展示等多方面支撑。"
        "在实现 InterviewAgent 的过程中，我进一步理解了如何将复杂任务拆解为多个可控节点，如何通过 RAG/GraphRAG 改善问题质量，"
        "以及如何通过测试和后台数据让系统更可解释、更可维护。"
    )

    add_heading(doc, "附录：截图插入建议", 1)
    add_para(
        doc,
        "请在 Word 中按下面建议补充截图。插图时建议统一宽度约 14-15 厘米，保持图片清晰，图注放在图片下方。"
    )
    for item in [
        "图 1：登录/注册页面。用于展示用户进入系统的第一视图。",
        "图 2：README 架构图或自己绘制的系统架构图。用于展示 FastAPI、JWT、Session Store、LangGraph、RAG、WebSocket 的关系。",
        "图 3：上传页面。建议包含简历上传、JD 输入、历史材料复用和技术面/HR 面选择。",
        "图 4：面试对话页面。建议截到阶段导航、候选人回答、面试官追问和输入区域。",
        "图 5：管理员后台页面。建议使用 http://127.0.0.1:8001/admin 登录后截图，展示覆盖领域和统计卡片。",
        "图 6：测试或接口页面。建议截图 pytest/verify_fast 通过结果，或 FastAPI /docs 接口列表。",
    ]:
        add_bullet(doc, item)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
