"""Generate the InterviewAgent software project management report set."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from scripts.project_reports_data import PROJECT, money, percent


OUTPUT_DIR = Path("output/project_management_reports")
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
RED = "9B1C1C"
GOLD = "7A5A00"
TOTAL_DXA = 9360
REPORTS = [
    ("01_项目建议书.docx", "项目建议书"),
    ("02_可行性分析报告.docx", "可行性分析报告"),
    ("03_项目计划书.docx", "项目计划书"),
    ("04_项目估算报告.docx", "项目估算报告"),
    ("05_项目进度和成本控制报告.docx", "项目进度和成本控制报告"),
    ("06_项目质量管理报告.docx", "项目质量管理报告"),
    ("07_项目风险管理报告.docx", "项目风险管理报告"),
    ("08_项目团队建设和管理报告.docx", "项目团队建设和管理报告"),
    ("09_项目总结报告.docx", "项目总结报告"),
]


def set_run_font(run, size=None, bold=None, color=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MID_GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color=MID_GRAY)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.1

    heading_specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.167


def new_document(report_title):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run(f"{PROJECT['short_name']} | {report_title}")
    set_run_font(run, size=9, color=MID_GRAY)
    add_page_field(section.footer.paragraphs[0])
    return doc


def add_cover(doc, report_title, subtitle="软件项目管理文件"):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(PROJECT["name"])
    set_run_font(r, size=25, bold=True, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(report_title)
    set_run_font(r, size=22, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(72)
    r = p.add_run(subtitle)
    set_run_font(r, size=12, color=MID_GRAY)

    for label, value in [
        ("作业编制小组", "王明哲、江悦铭、周士斌"),
        ("计划周期", "20 周"),
        ("文档版本", PROJECT["version"]),
        ("编制日期", PROJECT["report_date"]),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(f"{label}：{value}")
        set_run_font(r, size=11, color="344054")
    doc.add_page_break()


def add_document_control(doc, purpose, status="批准版"):
    doc.add_heading("文档控制", level=1)
    add_management_table(
        doc,
        ["属性", "内容"],
        [
            ["文档目的", purpose],
            ["编制", "王明哲、江悦铭、周士斌（软件项目管理作业小组）"],
            ["案例属性", "虚构软件项目管理案例，不代表现实项目执行进度"],
            ["版本与状态", f"{PROJECT['version']} / {status}"],
            ["适用范围", "InterviewAgent 建设、验收与收尾全过程"],
        ],
        [1900, 7460],
    )


def add_static_toc(doc, items):
    doc.add_heading("目录", level=1)
    for idx, item in enumerate(items, 1):
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(item)
        set_run_font(r)
    doc.add_page_break()


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.widow_control = True
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        set_run_font(first, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullets(doc, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.keep_together = True
        r = p.add_run(item)
        set_run_font(r)


def add_callout(doc, label, text, color=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_geometry(table, [TOTAL_DXA])
    cell = table.cell(0, 0)
    shade_cell(cell, LIGHT_GRAY)
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    r = p.add_run(f"{label}：")
    set_run_font(r, bold=True, color=color)
    r = p.add_run(text)
    set_run_font(r)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def set_table_geometry(table, widths):
    if sum(widths) != TOTAL_DXA:
        raise ValueError(f"Table widths must total {TOTAL_DXA}, got {sum(widths)}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TOTAL_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")


def add_management_table(doc, headers, rows, widths, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, value in enumerate(headers):
        cell = header.cells[idx]
        shade_cell(cell, LIGHT_BLUE)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(value))
        set_run_font(r, size=font_size, bold=True, color=DARK_BLUE)

    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if idx > 0 and len(str(value)) < 18:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    set_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_common_context(doc):
    doc.add_heading("项目概况", level=1)
    add_para(
        doc,
        "InterviewAgent 是面向求职者与人才评价场景的智能面试评估系统。系统接收简历和职位描述，组织多阶段技术或 HR 面试，通过 GraphRAG、题库检索和大模型生成追问，并支持文本、实时语音、会话管理及结构化评价报告。",
    )
    add_management_table(
        doc,
        ["维度", "基线"],
        [
            ["项目组织", "虚构团队 13 人：陈昊（项目经理）、赵宁（产品经理）、林泽宇（架构/AI 负责人）及 10 名开发、测试、UX、DevOps/安全成员"],
            ["周期", "规划总工期 20 周"],
            ["方法", PROJECT["method"]],
            ["预算基线", money(PROJECT["bac"])],
            ["目标版本", "可部署、可验收的 V1.0"],
            ["案例说明", PROJECT["scenario_note"]],
        ],
        [1700, 7660],
    )


def build_proposal():
    doc = new_document("项目建议书")
    add_cover(doc, "项目建议书", "立项决策与投资建议")
    add_document_control(doc, "说明建设必要性、目标、边界、投资与立项条件")
    add_static_toc(doc, ["执行摘要", "问题与机会", "拟建方案", "项目目标与范围", "收益与投资", "治理与实施建议", "立项结论"])
    doc.add_heading("执行摘要", level=1)
    add_callout(doc, "建议", "批准 InterviewAgent V1.0 立项，按 20 周总进度组织 13 人跨职能项目团队，预算基线为 %s。" % money(PROJECT["bac"]))
    add_para(doc, "传统模拟面试依赖导师时间，反馈口径不一致，难以形成可追溯的能力画像。InterviewAgent 将面试流程、追问依据和评分标准结构化，使用户能够重复练习并获得相对一致的诊断反馈。")
    add_common_context(doc)
    doc.add_heading("问题与机会", level=1)
    add_bullets(doc, [
        "求职者缺少低成本、可重复、与岗位要求匹配的面试练习环境。",
        "通用聊天机器人缺少阶段控制、岗位知识检索和统一评价量表，反馈可解释性不足。",
        "企业内部培训或高校就业辅导需要可规模化的模拟面试工具与过程数据。",
        "大模型、语音接口和混合检索技术已具备工程化组合条件，但必须通过质量门禁控制不确定性。",
    ])
    doc.add_heading("拟建方案", level=1)
    add_management_table(doc, ["能力域", "V1.0 建设内容", "价值"], [
        ["面试编排", "技术与 HR 多阶段流程、专项练习、阶段切换", "过程可控、可追踪"],
        ["内容理解", "简历解析、JD 技术点抽取、题库与上下文融合", "问题与候选人背景匹配"],
        ["GraphRAG", "知识图谱、BM25、向量检索和融合排序", "形成有层次的技术追问"],
        ["多模态", "文本、WebSocket、ASR、TTS", "接近真实面试交互"],
        ["评价报告", "维度评分、证据、弱项、参考答案、学习建议", "形成可解释反馈"],
        ["平台能力", "认证、会话隔离、历史记录、部署与日志", "支持多用户使用和运维"],
    ], [1500, 5160, 2700])
    doc.add_heading("项目目标与范围", level=1)
    add_bullets(doc, [
        "在第 11 周完成文本面试主流程演示，第 18 周达到系统验收条件。",
        "关键业务流程通过率不低于 95%，阶段流转正确率不低于 98%。",
        "评价报告结构化成功率不低于 99%，RAG Recall@5 不低于 0.80。",
        "严重缺陷清零，建立隐私、访问控制、降级和监控机制。",
    ])
    doc.add_heading("范围边界", level=2)
    add_management_table(doc, ["范围内", "范围外"], [[
        "Web 端面试、账号与会话、简历/JD、Agent、GraphRAG、语音、评价、部署文档",
        "原生移动端、企业 ATS 深度集成、正式招聘自动决策、私有化多租户商业平台",
    ]], [4680, 4680])
    doc.add_heading("收益与投资", level=1)
    add_para(doc, "预算基线由 %s 人力成本、%s 非人力直接成本和 %s 管理储备组成，总计 %s。收益以用户练习效率、服务可复制性、技术资产积累和后续产品化机会衡量，不承诺未经验证的直接现金回报。" % (money(PROJECT["labor_cost"]), money(sum(PROJECT["non_labor_costs"].values())), money(PROJECT["management_reserve"]), money(PROJECT["bac"])))
    add_management_table(doc, ["收益类别", "预期结果", "验证方式"], [
        ["用户价值", "完成一次结构化模拟面试并获得可解释报告", "可用性测试、完成率、满意度"],
        ["运营价值", "降低重复人工模拟面试的边际成本", "单次服务时长与调用成本"],
        ["技术资产", "形成 Agent 编排、RAG 评测与语音集成能力", "组件复用和评测记录"],
        ["管理价值", "形成可追踪范围、质量、风险和成本基线", "项目审计与收尾复盘"],
    ], [1800, 4560, 3000])
    doc.add_heading("治理与实施建议", level=1)
    add_bullets(doc, [
        "由虚构项目经理陈昊负责项目治理，范围、成本和关键技术决策实行变更审批。",
        "每两周交付一个可演示增量，在第 3、5、11、15、18、20 周设置阶段门。",
        "优先保证文本主流程和评价闭环，语音体验与扩展能力不得挤占验收必需范围。",
        "涉及真实简历和面试内容时执行最小化采集、访问隔离和保留期限控制。",
    ])
    doc.add_heading("立项结论", level=1)
    add_callout(doc, "结论", "项目具备明确需求和可实现路径，建议附条件立项：冻结 V1.0 范围、建立 AI 质量评测集、落实数据保护措施，并按阶段门释放管理储备。")
    return doc


def build_feasibility():
    doc = new_document("可行性分析报告")
    add_cover(doc, "可行性分析报告", "技术、经济、运营与合规论证")
    add_document_control(doc, "评价项目是否值得实施以及实施所需约束条件")
    add_static_toc(doc, ["分析结论", "方案比较", "技术可行性", "经济可行性", "运营与组织可行性", "进度可行性", "法律与合规", "敏感性分析", "结论"])
    doc.add_heading("分析结论", level=1)
    add_callout(doc, "总体判断", "在限定 V1.0 范围、采用成熟开源框架和第三方模型服务、配置质量与隐私控制的前提下，项目总体可行，推荐实施方案 B。")
    add_common_context(doc)
    doc.add_heading("备选方案比较", level=1)
    add_management_table(doc, ["方案", "说明", "优势", "局限", "加权评分"], [
        ["A 采购通用面试 SaaS", "配置现有产品", "上线快、运维少", "流程与 GraphRAG 难定制、数据受制于供应商", "72/100"],
        ["B 自研核心+托管模型", "自研平台、Agent、RAG，调用成熟模型/语音 API", "差异化强、工期可控、资产可沉淀", "存在供应商和调用成本风险", "86/100"],
        ["C 全栈私有化自研", "模型、语音、检索与平台均自建", "控制力最强", "即使 13 人团队也难在 20 周内完成模型训练与完整平台建设", "55/100"],
    ], [1100, 2300, 2100, 2760, 1100])
    doc.add_heading("技术可行性", level=1)
    add_para(doc, "FastAPI、LangGraph、Neo4j/NetworkX、ChromaDB、BM25、SentenceTransformer 与 WebSocket 均有成熟实现。最大技术不确定性不是单项技术可用性，而是多组件组合后的延迟、状态一致性、LLM 输出稳定性和检索质量。")
    add_management_table(doc, ["技术域", "成熟度", "验证重点", "结论"], [
        ["Agent 编排", "中高", "阶段状态、上下文隔离、异常恢复", "可行，需状态机测试"],
        ["GraphRAG", "中", "召回、排序、图谱质量、回退", "可行，需评测集驱动"],
        ["实时语音", "中", "弱网、首包延迟、打断和降级", "可行，是主要进度风险"],
        ["结构化评价", "中", "JSON 稳定性、证据一致性、偏差", "可行，需 Schema 与复核"],
        ["平台与部署", "高", "认证、会话隔离、配置和日志", "可行"],
    ], [1800, 1200, 3800, 2560])
    doc.add_heading("经济可行性", level=1)
    add_para(doc, "项目预算基线为 %s，其中人力占比 %s。对于验证差异化智能面试能力的 V1.0，自研核心与托管服务组合避免了训练模型和自建语音基础设施的高额投入。" % (money(PROJECT["bac"]), percent(PROJECT["labor_cost"] / PROJECT["bac"])))
    add_management_table(doc, ["成本项", "金额", "说明"], [["人力影子成本", money(PROJECT["labor_cost"]), "13 人按资源日历计划投入 8,440 小时"]] + [[k, money(v), "直接项目费用"] for k, v in PROJECT["non_labor_costs"].items()] + [["管理储备", money(PROJECT["management_reserve"]), "应对已知未知风险"]], [3500, 1800, 4060])
    doc.add_heading("运营与组织可行性", level=1)
    add_bullets(doc, [
        "13 人跨职能团队覆盖项目管理、产品、架构/AI、前后端、测试、UX 和 DevOps/安全，角色边界完整。",
        "系统架构师与 AI 技术负责人仍是关键资源，必须通过接口文档、代码评审和知识备份降低单点风险。",
        "系统应定位为辅助练习和能力诊断工具，不作为无人工复核的招聘淘汰决策系统。",
        "上线后至少需要模型调用监控、故障降级、数据删除请求和内容反馈处理机制。",
    ])
    doc.add_heading("进度可行性", level=1)
    add_para(doc, "自下而上估算得到 8,440 工时。13 名成员按阶段错峰投入，并非全部人员连续 20 周满负荷。计划通过资源日历、范围优先级、阶段门和管理储备控制偏差，而不通过持续加班弥补系统性问题。")
    add_management_table(doc, ["关键路径活动", "周次", "缓冲/控制"], [[p["name"], p["weeks"], "阶段验收后进入下一阶段"] for p in PROJECT["phases"]], [4100, 1600, 3660])
    doc.add_heading("法律、伦理与合规可行性", level=1)
    add_bullets(doc, [
        "简历、录音、转写和评价属于敏感度较高的个人信息，应取得明确同意并限定用途。",
        "建立最小权限、传输加密、口令安全、日志脱敏、数据保留期限和删除机制。",
        "对模型生成的评分与建议进行免责声明、证据展示和人工复核提示，防止误用。",
        "题库、简历和职位描述的来源需具备合法授权，避免存储或输出受保护内容。",
    ])
    doc.add_heading("敏感性分析", level=1)
    add_management_table(doc, ["情景", "影响", "应对后预计结果"], [
        ["模型/语音调用价格上升 50%", "直接成本增加约 ¥11,000", "仍可由部分管理储备覆盖"],
        ["语音集成延误 2 周", "关键路径延长", "降级为文本必选、语音限量发布，控制为 1 周延期"],
        ["RAG 指标未达 0.80", "技术面试质量不达标", "缩减知识域、提高人工题库权重后再验收"],
        ["核心成员可用率下降 20%", "关键模块延期 2-3 周", "知识转移、削减非核心范围并动用储备"],
    ], [3000, 2900, 3460])
    doc.add_heading("结论与约束条件", level=1)
    add_callout(doc, "结论", "选择方案 B：自研核心能力并使用托管模型与语音服务。立项条件包括范围冻结、评测集先行、隐私控制、文本降级路径和阶段性投资审查。")
    return doc


WBS_ROWS = [
    ["1.1", "项目章程与干系人识别", "章程、干系人登记册", "陈昊"],
    ["2.1", "用户场景与需求规格", "需求规格说明书", "赵宁"],
    ["2.2", "交互原型与验收标准", "可测试原型、验收清单", "蒋欣/赵宁"],
    ["3.1", "总体架构与接口契约", "架构说明、OpenAPI 契约", "林泽宇/刘晨"],
    ["3.2", "Agent、RAG 与语音技术验证", "技术验证记录", "林泽宇/孙博"],
    ["4.1", "认证、用户与会话管理", "后端服务与数据库", "刘晨/唐宇"],
    ["4.2", "多阶段 Agent 工作流", "可运行面试图", "郭嘉怡"],
    ["4.3", "GraphRAG 与题库检索", "混合检索与评测", "孙博"],
    ["4.4", "Web 面试工作台", "响应式前端", "许婧/何嘉"],
    ["4.5", "结构化评价与报告", "评价服务与展示", "郭嘉怡/刘晨"],
    ["5.1", "WebSocket、ASR 与 TTS", "语音交互链路", "唐宇/何嘉"],
    ["5.2", "端到端集成与降级", "集成版本", "全员"],
    ["6.1", "功能、接口与安全测试", "测试报告、缺陷清单", "郑凯/韩磊"],
    ["6.2", "RAG、Agent 与体验评测", "评测报告", "马琳/孙博"],
    ["7.1", "部署、监控与运维移交", "部署包、运行手册", "韩磊"],
    ["7.2", "验收、复盘与归档", "验收记录、总结报告", "陈昊"],
]


def build_plan():
    doc = new_document("项目计划书")
    add_cover(doc, "项目计划书", "范围、WBS、资源、进度、成本、风险与质量基线")
    add_document_control(doc, "规定项目执行和控制的正式基线")
    add_static_toc(doc, ["项目管理方法", "范围管理计划", "WBS", "资源计划", "进度计划", "成本计划", "质量计划", "风险计划", "沟通与变更", "验收与收尾"])
    add_common_context(doc)
    doc.add_heading("项目管理方法", level=1)
    add_para(doc, "项目采用阶段门控制下的两周迭代。预测型方法用于固定范围、里程碑、预算和验收门槛；迭代方法用于逐步验证 Agent 行为、RAG 质量和交互体验。每个迭代包含计划、开发、评审、测试和回顾。")
    add_management_table(doc, ["治理事件", "频率/节点", "输出"], [
        ["每日同步", "工作日 15 分钟", "阻塞项、当日目标"],
        ["迭代计划与评审", "每两周", "迭代目标、可演示增量、反馈"],
        ["风险与质量审查", "每周", "风险登记册、缺陷趋势、指标"],
        ["阶段门", "第 3/5/11/15/18/20 周", "批准、整改或调整范围"],
        ["变更控制", "按需，48 小时内评估", "变更请求、影响分析和决定"],
    ], [2700, 2400, 4260])
    doc.add_heading("范围管理计划", level=1)
    add_bullets(doc, [
        "范围基线由产品范围说明、WBS、WBS 字典和验收标准组成。",
        "需求使用 Must/Should/Could/Won't 分级；Must 项未完成时不得以新增 Could 项替代。",
        "任何影响里程碑、预算、数据边界或外部接口的变更必须书面评估。",
        "发起人批准基线变更；项目经理可在不改变验收目标的前提下调整任务顺序。",
    ])
    doc.add_heading("工作分解结构（WBS）", level=1)
    add_management_table(doc, ["WBS", "工作包", "主要交付物", "责任人"], WBS_ROWS, [900, 3000, 3560, 1900], 8.5)
    doc.add_heading("WBS 字典与完成定义", level=2)
    add_para(doc, "每个工作包必须具备明确输入、负责人、可检查交付物和完成标准。代码完成不等于工作包完成；只有通过代码评审、自动化测试、文档更新和演示验收后才能计入挣值。")
    doc.add_heading("资源计划", level=1)
    add_management_table(doc, ["成员", "职责", "计划工时", "影子费率", "主要备份安排"], [[name, info["role"], info["planned_hours"], f"¥{info['rate']}/小时", "关键模块至少一名成员完成评审和文档备份"] for name, info in PROJECT["team"].items()], [1200, 3000, 1200, 1500, 2460], 8.5)
    doc.add_heading("责任分配矩阵（RACI）", level=2)
    add_management_table(doc, ["交付域", "项目经理", "产品", "架构/AI", "后端", "前端", "QA/运维"], [
        ["需求与项目治理", "A", "R", "C", "C", "C", "I"],
        ["Agent 与 GraphRAG", "A", "C", "R", "C", "I", "C"],
        ["后端、数据与部署", "A", "I", "C", "R", "C", "C"],
        ["前端与交互体验", "A", "C", "C", "C", "R", "C"],
        ["语音链路", "A", "C", "C", "R", "R", "C"],
        ["质量与验收", "A", "C", "C", "C", "C", "R"],
    ], [2100, 1210, 1100, 1350, 1200, 1200, 1200], 8.0)
    doc.add_heading("进度计划", level=1)
    add_management_table(doc, ["阶段", "周次", "工时", "退出条件"], [[p["name"], p["weeks"], p["hours"], "阶段交付物评审通过"] for p in PROJECT["phases"]], [3100, 1400, 1200, 3660])
    doc.add_heading("里程碑计划", level=2)
    add_management_table(doc, ["编号", "里程碑", "计划时间"], PROJECT["milestones"], [1200, 5860, 2300])
    doc.add_heading("关键路径与进度控制", level=2)
    add_para(doc, "关键路径为需求基线→架构与技术验证→Agent/后端主流程→语音与系统集成→系统测试→验收部署。语音集成最多保留一周项目缓冲；若消耗超过阈值，执行文本优先和语音限量发布方案。")
    doc.add_heading("成本计划", level=1)
    add_management_table(doc, ["阶段", "计划人力成本", "直接非人力分摊", "阶段控制账户"], [[p["name"], money(p["labor_cost"]), money(round(sum(PROJECT["non_labor_costs"].values()) * p["hours"] / 8440)), p["code"]] for p in PROJECT["phases"]], [3300, 2000, 2300, 1760])
    add_callout(doc, "成本基线", "直接成本 %s，加管理储备 %s，预算基线 BAC 为 %s。管理储备由项目经理提出、发起人批准后使用。" % (money(PROJECT["direct_cost"]), money(PROJECT["management_reserve"]), money(PROJECT["bac"])))
    doc.add_heading("质量计划", level=1)
    add_management_table(doc, ["质量目标", "门槛", "测量频率", "责任人"], [[name, target, "每迭代/阶段门", "模块负责人，项目经理审核"] for name, target in PROJECT["quality_targets"]], [3500, 1900, 1900, 2060])
    add_bullets(doc, [
        "需求和架构采用同行评审；关键接口采用契约测试；核心状态转换采用自动化测试。",
        "严重缺陷阻断发布，高缺陷必须有明确处置决定，中低缺陷纳入技术债台账。",
        "AI 输出除功能测试外，使用固定评测集监测相关性、稳定性、重复率和结构化成功率。",
        "每次阶段门审查质量指标、未关闭缺陷、测试覆盖和已知限制。",
    ])
    doc.add_heading("风险计划", level=1)
    add_management_table(doc, ["编号", "风险", "等级", "主要响应", "责任人"], PROJECT["risks"], [800, 2800, 900, 3660, 1200], 8.3)
    add_para(doc, "风险每周复审。高风险触发定量影响分析和应急演练；风险暴露下降后仍保留观察期，不以“已采取措施”直接等同于关闭。")
    doc.add_heading("沟通与干系人管理", level=1)
    add_management_table(doc, ["对象", "信息需求", "方式", "频率", "责任人"], [
        ["项目发起人", "范围、里程碑、预算、重大风险", "状态报告/阶段评审", "双周及阶段门", "陈昊"],
        ["项目团队", "任务、接口、阻塞、质量", "每日同步/看板", "每日", "全员"],
        ["试用用户", "可用性、问题相关性、报告价值", "演示/访谈/测试", "第 11、15、18 周", "赵宁/蒋欣"],
        ["运维与数据责任方", "配置、监控、数据处理", "移交评审", "第 18-20 周", "韩磊"],
    ], [1700, 3000, 1900, 1460, 1300])
    doc.add_heading("变更与配置管理", level=1)
    add_bullets(doc, [
        "变更请求记录原因、范围、工时、成本、风险、质量和上线影响。",
        "需求、接口、数据模型、提示词和评测集均纳入版本控制；生产密钥不进入代码库。",
        "基线变更批准后同步更新 WBS、进度、预算、风险和验收标准。",
        "紧急缺陷可先处置后补录，但必须在下一工作日完成影响复盘。",
    ])
    doc.add_heading("验收与收尾计划", level=1)
    add_bullets(doc, [
        "Must 范围全部交付，关键业务流程、AI 质量和安全门槛达到约定值。",
        "部署包、配置清单、运行手册、测试报告、风险与技术债清单齐备。",
        "完成用户验收演示、问题关闭确认、知识移交和项目资料归档。",
        "总结计划与实际偏差，记录经验教训并释放项目资源。",
    ], numbered=True)
    return doc


def build_estimate():
    doc = new_document("项目估算报告")
    add_cover(doc, "项目估算报告", "资源、工期与成本估算")
    add_document_control(doc, "记录估算方法、假设、计算过程与结果")
    add_static_toc(doc, ["估算摘要", "估算基础", "资源估算", "工期估算", "成本估算", "三点估算", "敏感性与置信度", "结论"])
    add_common_context(doc)
    doc.add_heading("估算摘要", level=1)
    add_callout(doc, "结果", "虚构项目团队共 13 人，按阶段计划投入 8,440 工时；总工期 20 周；预算基线 %s；建议估算区间为 %s 至 %s。" % (money(PROJECT["bac"]), money(round(PROJECT["bac"] * 0.9)), money(round(PROJECT["bac"] * 1.15))))
    doc.add_heading("估算基础与假设", level=1)
    add_bullets(doc, [
        "13 名成员依据资源日历错峰投入，计划总工时 8,440 小时；并非所有岗位连续 20 周满负荷。",
        "大模型、ASR/TTS 使用托管服务，不包含训练基础模型或采购 GPU 集群。",
        "V1.0 仅建设 Web 端和单一部署环境，不包含原生移动端与企业级多租户。",
        "影子费率包含工资、福利、设备和一般管理分摊，用于反映真实机会成本。",
        "管理储备按直接成本的 12% 计算，不分配到工作包基线，使用时履行审批。",
    ])
    doc.add_heading("资源估算", level=1)
    add_management_table(doc, ["资源组", "人数", "计划工时", "能力要求", "估算依据"], [
        ["项目管理与产品", "2 人", "1,400 小时", "范围、需求、计划、验收与干系人管理", "WBS 自下而上"],
        ["架构与 AI/Agent", "3 人", "2,400 计划小时", "架构、LangGraph、GraphRAG、评测", "关键路径与专家判断"],
        ["后端开发", "2 人", "1,520 小时", "FastAPI、数据库、认证、WebSocket", "WBS 自下而上"],
        ["前端开发", "2 人", "1,360 小时", "Web、流式交互、音频与报告展示", "WBS 自下而上"],
        ["测试与质量", "2 人", "1,040 小时", "自动化、性能、安全与 AI 质量", "测试工作量模型"],
        ["UX 与 DevOps/安全", "2 人", "720 小时", "原型、部署、监控与安全", "阶段投入估算"],
        ["云与模型资源", "1 套", "开发/测试/试运行", "LLM、ASR/TTS、数据库、日志", "参数估算"],
    ], [2100, 1100, 1500, 3100, 1560], 8.3)
    doc.add_heading("工期估算", level=1)
    add_para(doc, "工期采用 WBS 自下而上估算并结合依赖关系形成网络计划。总工时不能简单除以 13 人，因为产品、架构、开发、测试和部署具有先后依赖，架构与 AI 专家还是受约束资源，因此必须依据关键路径和资源日历确定 20 周基线。")
    add_management_table(doc, ["阶段", "工时", "理论人周", "计划周次", "并行性判断"], [[p["name"], p["hours"], f"{p['hours']/40:.1f}", p["weeks"], "部分并行，阶段门约束"] for p in PROJECT["phases"]], [3000, 1100, 1300, 1400, 2560])
    doc.add_heading("PERT 三点估算", level=2)
    pert_rows = [
        ["需求与原型", 1.5, 2.0, 3.0],
        ["架构与技术验证", 1.5, 2.0, 3.5],
        ["核心功能开发", 5.0, 6.0, 8.0],
        ["语音与系统集成", 3.0, 4.0, 6.0],
        ["测试、验收与部署", 4.0, 6.0, 8.0],
    ]
    calculated = []
    total = 0
    for name, o, m, p in pert_rows:
        e = (o + 4 * m + p) / 6
        total += e
        calculated.append([name, o, m, p, f"{e:.2f}"])
    add_management_table(doc, ["活动", "乐观", "最可能", "悲观", "期望周数"], calculated, [3000, 1300, 1500, 1300, 2260])
    add_para(doc, f"PERT 独立阶段期望值合计约 {total:.2f} 周。考虑部分活动并行和一周项目缓冲后，20 周基线具有可执行性，但对语音和端到端稳定性偏差较敏感。")
    doc.add_heading("成本估算", level=1)
    add_management_table(doc, ["阶段", "工时", "人力成本", "平均小时成本"], [[p["name"], p["hours"], money(p["labor_cost"]), f"¥{p['labor_cost']/p['hours']:.2f}"] for p in PROJECT["phases"]], [3900, 1300, 2200, 1960])
    add_management_table(doc, ["预算构成", "金额", "占 BAC"], [
        ["直接人力", money(PROJECT["labor_cost"]), percent(PROJECT["labor_cost"] / PROJECT["bac"])],
        ["非人力直接成本", money(sum(PROJECT["non_labor_costs"].values())), percent(sum(PROJECT["non_labor_costs"].values()) / PROJECT["bac"])],
        ["管理储备", money(PROJECT["management_reserve"]), percent(PROJECT["management_reserve"] / PROJECT["bac"])],
        ["预算基线 BAC", money(PROJECT["bac"]), "100.0%"],
    ], [4800, 2300, 2260])
    doc.add_heading("敏感性与估算置信度", level=1)
    add_management_table(doc, ["变量", "变化", "预算/工期影响", "敏感度"], [
        ["关键架构/AI 人员生产率", "-15%", "工期增加约 2-3 周或削减范围", "高"],
        ["模型与语音单价", "+50%", "成本增加约 ¥11,000", "中"],
        ["返工率", "由 10% 升至 20%", "增加约 160-240 小时", "高"],
        ["RAG 数据准备", "+1 周", "可能挤压核心开发或验收缓冲", "中高"],
    ], [2600, 1800, 3460, 1500])
    add_callout(doc, "估算结论", "20 周和 %s 作为批准基线；管理层应将 21-23 周、%s-%s 视为合理风险区间，并优先通过范围排序而非无计划加班应对偏差。" % (money(PROJECT["bac"]), money(round(PROJECT["bac"] * 1.02)), money(round(PROJECT["bac"] * 1.15))))
    return doc


def build_control():
    doc = new_document("项目进度和成本控制报告")
    add_cover(doc, "项目进度和成本控制报告", "20 周总进度安排与成本控制方案")
    add_document_control(doc, "规定项目总进度、成本基线、测量规则和偏差控制流程", "计划基线版")
    add_static_toc(doc, ["控制目标", "总进度安排", "关键路径与里程碑", "成本基线", "进度测量方法", "挣值管理规则", "偏差阈值与纠偏", "变更与报告机制"])
    add_common_context(doc)
    doc.add_heading("控制目标与基准", level=1)
    add_callout(doc, "控制基准", "本报告给出虚构项目从第 1 周至第 20 周的完整计划，不报告任何现实进展。进度以批准的 WBS、里程碑和资源日历为基准，成本以 BAC %s 为上限基准。" % money(PROJECT["bac"]))
    add_bullets(doc, [
        "保证关键路径工作包按顺序形成可验收交付物。",
        "每周更新工作包剩余工时，每两周在迭代评审中确认挣值。",
        "将范围、工期、资源和成本变化放入同一变更流程，防止局部优化破坏总计划。",
        "预先规定黄色和红色阈值，使项目执行后的偏差处理有客观依据。",
    ])
    doc.add_heading("20 周总进度安排", level=1)
    add_management_table(doc, ["阶段", "计划周次", "工时", "主要活动", "阶段出口"], [
        ["立项与启动", "1", "320", "章程、干系人、治理机制", "章程批准"],
        ["需求分析与原型", "2-3", "760", "场景、需求、原型、验收标准", "需求基线"],
        ["架构设计与验证", "4-5", "920", "架构、接口、Agent/RAG/语音验证", "技术评审"],
        ["核心功能开发", "6-11", "3,340", "平台、Agent、RAG、前端、评价", "文本主流程版本"],
        ["语音与系统集成", "12-15", "1,280", "ASR/TTS、WebSocket、端到端集成", "集成版本"],
        ["测试与验收准备", "16-18", "1,280", "功能、性能、安全、AI 评测、UAT", "验收评审"],
        ["部署移交与收尾", "19-20", "540", "部署、培训、移交、归档", "收尾评审"],
    ], [2100, 1200, 1100, 3060, 1900], 8.2)
    doc.add_heading("关键路径与里程碑", level=1)
    add_para(doc, "关键路径为需求基线→架构与接口契约→Agent/后端主流程→语音集成→系统测试→验收部署。项目经理每周检查关键路径活动的剩余工期和资源可用性；非关键工作仅可在不消耗关键资源的前提下并行。")
    add_management_table(doc, ["编号", "里程碑", "计划时间", "验收责任"], [[code, name, week, "项目经理组织、发起人/业务代表批准"] for code, name, week in PROJECT["milestones"]], [1000, 4300, 1700, 2360])
    doc.add_heading("成本基线与阶段预算", level=1)
    add_management_table(doc, ["控制账户", "阶段", "计划工时", "人力预算", "非人力预算分摊"], [[p["code"], p["name"], p["hours"], money(p["labor_cost"]), money(round(sum(PROJECT["non_labor_costs"].values()) * p["hours"] / 8440))] for p in PROJECT["phases"]], [1000, 2800, 1400, 2100, 2060], 8.2)
    add_para(doc, "直接成本为 %s，管理储备为 %s，BAC 为 %s。工作包负责人只能使用已分配的控制账户预算；管理储备不预先摊入工作包，动用时必须提交风险或变更依据。" % (money(PROJECT["direct_cost"]), money(PROJECT["management_reserve"]), money(PROJECT["bac"])))
    doc.add_heading("进度测量与完成规则", level=1)
    add_management_table(doc, ["工作类型", "推荐测量方法", "计入完成的条件"], [
        ["短期离散任务", "0/100 法", "交付物通过评审后计 100%"],
        ["跨两周工作包", "50/50 法", "正式启动计 50%，验收计剩余 50%"],
        ["长周期开发包", "加权里程碑", "设计、编码、测试、验收按预设权重"],
        ["持续支持活动", "按投入比例", "仅限项目管理、环境维护等持续工作"],
    ], [2200, 2400, 4760])
    doc.add_heading("挣值管理规则", level=1)
    add_management_table(doc, ["指标", "公式", "用途"], [
        ["PV", "状态周期应完成工作的预算价值", "与基线比较计划要求"],
        ["EV", "按完成规则确认的预算价值", "衡量客观完成量"],
        ["AC", "执行期累计发生成本", "反映资源消耗"],
        ["SV", "EV - PV", "小于 0 表示进度偏后"],
        ["CV", "EV - AC", "小于 0 表示成本效率不利"],
        ["SPI", "EV / PV", "进度效率指数"],
        ["CPI", "EV / AC", "成本效率指数"],
        ["EAC", "BAC / CPI 或 BAC / (CPI×SPI)", "根据偏差性质预测完工成本"],
    ], [1300, 3600, 4460])
    doc.add_heading("偏差阈值与纠偏流程", level=1)
    add_management_table(doc, ["等级", "进度/成本触发条件", "管理动作"], [
        ["绿色", "SPI、CPI 均 >=0.95，关键里程碑无滑移", "工作包负责人自行纠偏并周报"],
        ["黄色", "任一指数 0.90-0.95，或关键路径预计滑移 <=5 个工作日", "项目经理组织根因分析、资源调整和恢复计划"],
        ["红色", "任一指数 <0.90，或关键里程碑预计滑移 >5 个工作日，或预计超 BAC", "提交变更控制委员会，调整范围、预算或总工期"],
    ], [1300, 3860, 4200])
    add_bullets(doc, [
        "先确认偏差是否源于错误的完成度申报、范围变化或资源瓶颈。",
        "比较赶工、快速跟进、资源替换、范围降级等方案的成本与新增风险。",
        "选择纠偏方案后更新剩余工作预测，但不得直接覆盖原始基线。",
        "需要改变里程碑、BAC 或验收范围时，必须形成正式变更单。",
    ], numbered=True)
    doc.add_heading("变更与报告机制", level=1)
    add_management_table(doc, ["报告", "频率", "主要内容", "接收方"], [
        ["工作包进度更新", "每周", "完成量、剩余工时、阻塞和预测", "项目经理"],
        ["迭代控制报告", "每两周", "里程碑、PV/EV/AC、SPI/CPI、风险", "项目团队与发起人"],
        ["阶段门报告", "关键节点", "交付物、预算消耗、质量门槛、变更建议", "指导委员会"],
        ["例外报告", "触发红色阈值时", "根因、影响、备选方案和决策请求", "变更控制委员会"],
    ], [2500, 1600, 3660, 1600])
    add_callout(doc, "控制结论", "项目执行期间应持续保留原始 20 周进度基线和 %s 成本基线，所有状态数据在实施后按本报告规则采集；本作业不假定项目已经开始或完成。" % money(PROJECT["bac"]), color=GOLD)
    return doc


def build_quality():
    doc = new_document("项目质量管理报告")
    add_cover(doc, "项目质量管理报告", "质量规划、保证、控制与改进")
    add_document_control(doc, "说明质量体系、执行记录、度量结果和改进措施")
    add_static_toc(doc, ["质量方针", "质量目标", "质量保证", "质量控制", "AI 专项质量", "缺陷分析", "改进与结论"])
    add_common_context(doc)
    doc.add_heading("质量方针与原则", level=1)
    add_para(doc, "项目以可用、可解释、可恢复和保护用户数据为核心质量属性。对于生成式 AI，质量不能仅以“能返回内容”判断，而要同时检查流程正确性、检索依据、结构稳定性、延迟和失败降级。")
    doc.add_heading("质量目标与度量计划", level=1)
    quality_plan = [
        ["关键业务流程通过率", ">=95%", "端到端测试套件", "阶段门"],
        ["阶段流转正确率", ">=98%", "状态转换自动化用例", "每次迭代"],
        ["评价报告结构化成功率", ">=99%", "固定样本批量验证", "模型/提示词变更时"],
        ["RAG Recall@5", ">=0.80", "标注评测集", "每次索引或检索策略变更"],
        ["严重缺陷遗留数", "0", "缺陷管理系统", "发布评审"],
        ["API P95（非 LLM）", "<800 ms", "性能测试与监控", "集成测试/试运行"],
        ["语音首包 P95", "<2.5 s", "弱网与并发场景测试", "语音阶段门"],
    ]
    add_management_table(doc, ["指标", "目标", "测量方法", "测量时点"], quality_plan, [3000, 1600, 3000, 1760])
    doc.add_heading("质量保证活动", level=1)
    add_management_table(doc, ["活动", "对象", "频率", "证据"], [
        ["需求与验收评审", "Must/Should 范围、验收标准", "基线及变更时", "评审记录、变更单"],
        ["架构与接口评审", "状态模型、API、数据与降级", "关键设计前", "决策记录、接口契约"],
        ["代码评审", "核心逻辑和安全边界", "每次合并", "评审意见、检查清单"],
        ["质量审计", "测试、缺陷、文档、配置", "每个阶段门", "质量门报告"],
        ["评测集回归", "Agent、RAG、报告结构", "每次模型/提示词变更", "指标趋势"],
    ], [2400, 3000, 1700, 2260])
    doc.add_heading("质量控制与测试策略", level=1)
    add_management_table(doc, ["测试层级", "范围", "完成标准"], [
        ["单元测试", "状态转换、权限、解析、评分工具函数", "关键逻辑和边界用例通过"],
        ["接口/契约测试", "认证、上传、会话、流式事件、报告", "契约无破坏性变化"],
        ["集成测试", "数据库、Redis、图谱、向量库、模型与语音", "降级路径和超时行为正确"],
        ["端到端测试", "注册到完成面试并查看报告", "核心场景通过率 >=95%"],
        ["安全与隐私测试", "越权、上传、密钥、日志、数据删除", "严重问题为零"],
        ["可用性测试", "面试工作台、语音状态、报告理解", "关键任务无阻断"],
    ], [2200, 4300, 2860])
    doc.add_heading("AI 与 RAG 专项质量", level=1)
    add_bullets(doc, [
        "固定模型版本、温度和提示词版本，避免同一评测无法复现。",
        "RAG 使用 Recall@K、MRR、问题相关性和无依据追问率，不以单个案例判断。",
        "评价报告验证 JSON Schema、证据引用、评分范围、弱项与建议一致性。",
        "对空检索、模型超时、内容安全拒答和语音失败设置可观察的降级结果。",
        "人工抽检重点关注偏见、幻觉、重复提问和将推断当作事实的情况。",
    ])
    doc.add_heading("缺陷分级与根因分析机制", level=1)
    add_management_table(doc, ["严重度", "定义", "修复时限", "发布规则"], [
        ["S1 阻断", "服务不可用、越权或数据泄露", "立即响应", "未关闭不得发布"],
        ["S2 严重", "核心流程失败、评分或数据明显错误", "当前迭代", "原则上不得遗留"],
        ["S3 一般", "非核心功能或兼容性问题", "排入最近迭代", "需评估并记录"],
        ["S4 轻微", "样式、文案和低影响体验问题", "按优先级安排", "可作为技术债"],
    ], [1700, 4100, 1700, 1860])
    add_para(doc, "每两周使用缺陷类别、注入阶段和逃逸阶段进行帕累托分析。重复出现的状态、接口、第三方服务和模型格式问题必须开展 5Why 或鱼骨分析，并将根因措施转化为检查表、自动化测试或设计规则。")
    doc.add_heading("持续改进计划", level=1)
    add_bullets(doc, [
        "将流式事件与评价 Schema 纳入自动契约测试，变更前先运行兼容性检查。",
        "将弱网、超时、空检索和模型格式漂移用例加入每次回归。",
        "把 RAG 评测结果与提示词、语料和模型版本绑定，形成可追溯基线。",
        "试运行期间持续监控语音首包延迟；未达到 2.5 秒目标时不得将语音能力标记为全面验收。",
    ], numbered=True)
    add_callout(doc, "质量结论", "项目应以量化门槛、独立测试证据和阶段门评审决定是否进入下一阶段。本报告只定义质量计划和控制规则，不预设项目已经达到验收标准。")
    return doc


def build_risk():
    doc = new_document("项目风险管理报告")
    add_cover(doc, "项目风险管理报告", "识别、分析、响应与残余风险")
    add_document_control(doc, "记录风险方法、风险登记册、响应执行和残余暴露")
    add_static_toc(doc, ["风险方法", "风险概况", "风险登记册", "应急与触发器", "风险复审安排", "风险接受与结论"])
    add_common_context(doc)
    doc.add_heading("风险管理方法", level=1)
    add_para(doc, "风险以概率 1-5、影响 1-5 评分，暴露值为两者乘积。15-25 为高风险，8-14 为中风险，1-7 为低风险。风险责任人每周复审触发器和响应状态；问题一旦发生即转入问题日志，不继续以未发生风险管理。")
    add_management_table(doc, ["影响维度", "低", "中", "高"], [
        ["进度", "<3 天", "3-10 天", ">10 天或关键里程碑失守"],
        ["成本", "<2% BAC", "2%-8% BAC", ">8% BAC"],
        ["质量/安全", "局部体验", "核心功能降级", "数据泄露、严重错误或无法验收"],
    ], [2200, 2200, 2200, 2760])
    doc.add_heading("风险登记册", level=1)
    risk_rows = [
        ["R01", "LLM 格式漂移", "4", "4", "16 高", "Schema、重试、模板降级", "主动降低", "林泽宇"],
        ["R02", "GraphRAG 相关性不足", "4", "4", "16 高", "评测集、混合检索、回退", "主动降低", "孙博"],
        ["R03", "第三方接口中断", "3", "5", "15 高", "超时、熔断、文本降级", "降低/转移", "刘晨"],
        ["R04", "个人信息泄露", "3", "5", "15 高", "最小采集、隔离、脱敏、删除", "规避/降低", "韩磊"],
        ["R05", "语音延迟", "4", "3", "12 中", "流式、弱网测试、状态提示", "主动降低", "何嘉"],
        ["R06", "核心 AI 人员单点依赖", "4", "4", "16 高", "文档、评审、结对和备份", "主动降低", "林泽宇"],
        ["R07", "范围蔓延", "4", "3", "12 中", "基线、MoSCoW、变更控制", "规避", "陈昊"],
        ["R08", "接口反复变更", "4", "3", "12 中", "契约优先、Mock、版本化", "主动降低", "刘晨"],
    ]
    add_management_table(doc, ["ID", "风险", "P", "I", "暴露", "响应", "状态", "责任人"], risk_rows, [650, 1650, 500, 500, 850, 2700, 1300, 1210], 7.8)
    doc.add_heading("触发器与应急计划", level=1)
    add_management_table(doc, ["风险", "量化触发器", "应急措施"], [
        ["R01", "结构化失败率连续两天 >1%", "切换稳定模型/模板，暂停提示词发布"],
        ["R02", "Recall@5 <0.80 或无关追问率 >10%", "缩小知识域，提高题库和关键词权重"],
        ["R03", "5 分钟错误率 >10%", "熔断第三方接口，切换文本或缓存回答"],
        ["R04", "出现越权读取、密钥泄露或敏感日志", "停止服务、隔离、审计、通知和修复"],
        ["R05", "语音首包 P95 >3 秒", "显示状态并自动提供文本交互"],
        ["R06", "关键负责人连续一周负荷 >110%", "调整范围，授权协调，安排结对"],
    ], [1200, 3600, 4560])
    doc.add_heading("风险复审与问题转化安排", level=1)
    add_para(doc, "风险登记册在项目启动时建立，此后每周复审概率、影响、触发器、响应责任和次生风险。某一不确定事件一旦发生，应转入问题日志并指定解决时限；风险登记册保留其来源和响应效果，不以主观判断直接关闭。")
    add_management_table(doc, ["评审节点", "重点复审内容", "预期输出"], [
        ["第1周启动评审", "范围、组织、供应商和数据边界风险", "初始风险登记册"],
        ["第5周技术评审", "Agent、RAG、语音与架构风险", "技术响应和储备建议"],
        ["第11周开发阶段门", "接口、质量、范围和关键资源风险", "集成前风险清单"],
        ["第15周集成阶段门", "性能、第三方服务、数据与兼容风险", "测试优先级和应急演练"],
        ["第18周验收评审", "发布、运维和残余风险", "风险接受/延期决定"],
        ["第20周收尾评审", "运营移交风险与责任", "移交登记册"],
    ], [2200, 4300, 2860])
    doc.add_heading("风险接受与结论", level=1)
    add_bullets(doc, [
        "第三方模型和语音服务可用性仍依赖供应商，应保持文本降级与配额监控。",
        "AI 评价可能存在偏差和不可解释推断，系统不得作为自动招聘决策依据。",
        "真实用户数据进入系统后，数据保留、删除和审计能力需持续运营。",
        "核心 AI 知识可能集中于少数人员，收尾前必须完成 GraphRAG 与 Agent 运维手册及恢复演练。",
    ])
    add_callout(doc, "风险结论", "高风险不得仅凭责任人承诺接受；必须有量化触发器、经验证的响应措施和批准人。项目进入验收前，应重新评估残余风险并决定接受、延期或调整范围。")
    return doc


def build_team():
    doc = new_document("项目团队建设和管理报告")
    add_cover(doc, "项目团队建设和管理报告", "组织、协作、绩效与团队发展")
    add_document_control(doc, "说明虚构项目团队的组织设计、建设方案和管理机制")
    add_static_toc(doc, ["团队目标与结构", "职责与授权", "技能与负荷", "工作协议", "沟通与冲突", "绩效与发展", "团队复盘"])
    add_common_context(doc)
    doc.add_heading("团队目标与组织结构", level=1)
    add_para(doc, "虚构项目采用项目型组织，由项目经理陈昊对范围、进度、成本和跨职能协调负责。产品、架构/AI、前后端、测试、UX 与 DevOps/安全分别设置明确负责人；重大范围、预算、数据安全和验收决定提交项目指导委员会。")
    add_management_table(doc, ["成员", "岗位", "主要职责", "计划投入"], [[name, info["role"], {
        "陈昊": "项目治理、计划、预算、风险、干系人与变更",
        "赵宁": "用户研究、需求、产品范围和验收标准",
        "林泽宇": "架构、AI 技术路线、接口原则和技术评审",
        "孙博": "GraphRAG、知识图谱、检索与评测",
        "郭嘉怡": "Agent 工作流、提示词、评价与算法验证",
        "刘晨": "API、认证、会话、数据模型与后端集成",
        "唐宇": "语音后端、WebSocket、第三方服务与数据处理",
        "许婧": "前端架构、面试工作台和报告展示",
        "何嘉": "流式交互、音频前端和浏览器兼容",
        "郑凯": "功能、接口、自动化和性能测试",
        "马琳": "质量保证、AI 评测、缺陷与验收管理",
        "蒋欣": "交互原型、视觉规范和可用性测试",
        "韩磊": "CI/CD、环境、监控、配置与安全",
    }[name], f"{info['planned_hours']} 小时"] for name, info in PROJECT["team"].items()], [1300, 2300, 4060, 1700], 8.0)
    doc.add_heading("技能矩阵与知识备份", level=1)
    add_management_table(doc, ["关键能力", "主责", "备份/评审", "知识保全措施"], [
        ["项目计划与变更", "陈昊", "赵宁", "基线、周报、决策日志和变更台账"],
        ["架构与 Agent", "林泽宇/郭嘉怡", "孙博/刘晨", "架构决策记录、提示词版本和结对评审"],
        ["GraphRAG 与评测", "孙博", "林泽宇/马琳", "数据字典、评测集、索引和回退手册"],
        ["后端与数据库", "刘晨/唐宇", "韩磊", "OpenAPI、迁移脚本和运行手册"],
        ["前端与语音交互", "许婧/何嘉", "蒋欣/郑凯", "组件文档、浏览器与弱网兼容清单"],
        ["质量、安全与部署", "马琳/韩磊", "郑凯/刘晨", "门禁规则、应急预案和恢复演练"],
    ], [2200, 2200, 2200, 2760])
    doc.add_heading("团队工作协议", level=1)
    add_bullets(doc, [
        "每日同步只讨论目标、完成、阻塞和协助，不替代技术专题会议。",
        "接口变化先更新契约和示例，再修改实现；破坏性变化必须获得消费方确认。",
        "关键代码至少一人评审；安全、状态机和数据边界必须由非作者复核。",
        "阻塞超过 4 小时主动提出，超过 1 个工作日由项目经理协调范围或资源。",
        "争议先基于用户目标、验收标准和数据判断；技术域负责人提出方案，跨域影响由项目经理裁决并记录。",
        "可持续工作优先，连续高负荷触发任务调整，不把长期加班作为进度策略。",
    ])
    doc.add_heading("沟通机制", level=1)
    add_management_table(doc, ["机制", "参与者", "节奏", "目的"], [
        ["每日同步", "全员", "每日", "发现阻塞与依赖"],
        ["技术设计评审", "相关负责人+一名评审者", "按需", "形成可追踪决策"],
        ["迭代评审/回顾", "全员及试用代表", "每两周", "验证增量并改进协作"],
        ["一对一负荷检查", "项目经理与成员", "双周", "识别压力、成长和支持需求"],
        ["风险/质量会", "全员", "每周", "更新风险、缺陷和指标"],
    ], [2300, 2500, 1500, 3060])
    doc.add_heading("负荷与单点管理", level=1)
    add_para(doc, "林泽宇、孙博等架构和 AI 人员位于关键路径，是主要受约束资源；测试与 DevOps 则在中后期集中投入。项目经理应通过资源直方图和未来四周滚动计划提前识别超过 100% 的分配，避免依靠长期加班。")
    add_management_table(doc, ["阶段", "主要投入角色", "负荷风险", "管理动作"], [
        ["1-5 周", "项目、产品、UX、架构/AI、技术负责人", "专家并行参与多个验证", "限制技术验证数量，明确退出标准"],
        ["6-11 周", "AI、后端、前端、测试", "接口依赖和开发 WIP 过高", "契约冻结、限制 WIP、测试前移"],
        ["12-15 周", "后端、前端、AI、测试、DevOps", "语音和端到端集成冲突", "联合攻关，暂停低优先级需求"],
        ["16-18 周", "测试、开发、产品、安全", "缺陷修复争抢开发资源", "严重度分诊和每日缺陷会"],
        ["19-20 周", "DevOps、项目、产品、技术负责人", "移交文档和部署集中", "提前准备清单并执行恢复演练"],
    ], [1800, 3000, 2300, 2260])
    doc.add_heading("冲突与决策管理", level=1)
    add_management_table(doc, ["典型冲突", "处理方式", "结果"], [
        ["语音体验优化与按期验收冲突", "依据 Must 范围，设计文本降级并延后非关键优化", "以基线和验收优先级做决策"],
        ["流式事件字段由谁定义", "以后端契约为基线，前端参与用例评审", "减少重复联调"],
        ["RAG 指标与内容丰富度冲突", "先缩小知识域保证相关性，再逐步扩展", "以 Recall@5 >=0.80 作为准入门槛"],
    ], [3200, 3600, 2560])
    doc.add_heading("绩效、认可与发展", level=1)
    add_bullets(doc, [
        "绩效以团队目标和可验证交付物为主，不按代码行数或单纯加班时长评价。",
        "个人目标包含主责交付、跨域协作、质量结果、知识共享和风险主动暴露。",
        "迭代评审中公开认可解决复杂问题、帮助他人和提前发现风险的行为。",
        "项目经理发展授权与控制能力；技术成员发展跨域评审、质量和运维意识；测试和产品成员提升 AI 评测能力。",
    ])
    doc.add_heading("团队发展与复盘安排", level=1)
    add_para(doc, "团队建设按形成期、磨合期、规范期和高效期设计。启动阶段完成角色澄清和工作协议；开发早期重点解决接口边界与决策冲突；集成阶段通过联合问题小组形成跨职能协作；每两周回顾团队负荷、沟通质量和改进行动。")
    add_callout(doc, "团队结论", "13 人项目型组织能够覆盖 V1.0 所需专业能力。团队管理重点不是把编制作业的三名同学映射为开发人员，而是为虚构项目设置清晰岗位、责任、授权、绩效和知识备份机制。")
    return doc


def build_summary():
    doc = new_document("项目总结报告")
    add_cover(doc, "项目总结报告", "项目管理方案、基线与预期收尾总结")
    add_document_control(doc, "总结虚构项目的管理方案、计划基线、控制重点和预期收尾要求")
    add_static_toc(doc, ["总结说明", "项目目标与范围", "管理基线", "质量与风险体系", "组织与治理", "管理重点", "预期收尾", "总结结论"])
    add_common_context(doc)
    doc.add_heading("总结说明", level=1)
    add_callout(doc, "文档边界", "本报告总结的是软件项目管理作业中建立的虚构项目方案，不代表 InterviewAgent 已按该计划实施、验收或收尾。任何进度、成本和质量数值均为计划基线或控制目标。")
    doc.add_heading("项目目标与范围总结", level=1)
    add_management_table(doc, ["管理对象", "规划内容", "验收依据"], [
        ["多阶段面试流程", "技术/HR 流程、专项练习和阶段控制", "核心场景端到端用例"],
        ["GraphRAG 技术追问", "知识图谱、关键词、向量与融合检索", "Recall@5 与人工相关性评审"],
        ["文本与实时语音", "SSE/WebSocket、ASR/TTS 和文本降级", "性能、弱网和兼容性测试"],
        ["结构化评价报告", "评分、证据、弱项和学习建议", "Schema 稳定性与内容抽检"],
        ["平台与数据治理", "认证、会话隔离、历史、日志和部署", "安全、隐私和运行检查表"],
    ], [2600, 4200, 2560])
    doc.add_heading("管理基线总结", level=1)
    add_management_table(doc, ["基线", "批准值", "控制说明"], [
        ["范围", "V1.0 Must 交付物与 WBS", "范围变化必须走变更控制"],
        ["总工期", "20 周", "七阶段、七里程碑、关键路径控制"],
        ["计划工时", "8,440 小时", "13 人按资源日历错峰投入"],
        ["预算", money(PROJECT["bac"]), "含直接成本和 12% 管理储备"],
        ["开发方法", PROJECT["method"], "阶段门固定基线，迭代验证增量"],
    ], [2200, 2900, 4260])
    doc.add_heading("质量与风险体系总结", level=1)
    add_para(doc, "质量体系覆盖需求评审、架构评审、代码评审、自动化测试、AI/RAG 固定评测集、性能安全测试和用户验收。风险体系使用概率影响矩阵、责任人、触发器、预防措施、应急措施和残余风险审批。")
    add_management_table(doc, ["管理域", "关键门槛/规则", "决策用途"], [
        ["核心流程", "通过率 >=95%，严重缺陷为 0", "是否进入发布评审"],
        ["Agent 与报告", "阶段流转 >=98%，结构成功率 >=99%", "是否允许模型/提示词版本升级"],
        ["GraphRAG", "Recall@5 >=0.80", "知识域是否准入"],
        ["进度成本", "SPI/CPI 绿色 >=0.95，红色 <0.90", "是否启动纠偏或基线变更"],
        ["高风险", "必须有责任人、触发器和验证过的响应", "是否允许进入下一阶段"],
    ], [2200, 4200, 2960])
    doc.add_heading("组织与治理总结", level=1)
    add_para(doc, "虚构项目团队由 13 人组成，覆盖项目管理、产品、架构/AI、前后端、测试、UX 和 DevOps/安全。项目经理陈昊负责综合管理，技术域负责人拥有专业决策权，项目指导委员会负责重大范围、预算、数据风险和阶段门批准。")
    add_management_table(doc, ["治理层级", "主要职责", "核心输出"], [
        ["指导委员会/发起人", "批准章程、基线、重大变更和验收", "决策与阶段门批准"],
        ["项目经理", "综合计划、协调、控制、风险和沟通", "主计划、状态模板、变更台账"],
        ["产品与技术负责人", "需求、架构、交付和技术质量", "需求基线、设计与交付物"],
        ["质量与安全角色", "独立验证、门禁、缺陷和风险审查", "测试证据、质量与安全报告"],
    ], [2500, 4200, 2660])
    doc.add_heading("项目管理重点", level=1)
    add_management_table(doc, ["主题", "管理重点", "预防性安排"], [
        ["范围", "避免语音、移动端和知识域无限扩张", "MoSCoW、WBS 和变更委员会"],
        ["架构", "控制多组件接口与状态复杂度", "技术验证、契约优先和降级设计"],
        ["AI 质量", "防止主观演示替代量化验证", "固定评测集、版本绑定和人工抽检"],
        ["进度", "关键专家与集成活动受资源约束", "关键路径、资源日历和滚动四周计划"],
        ["成本", "模型调用和返工可能消耗储备", "控制账户、EVM 和储备审批"],
        ["团队", "跨职能目标与知识集中风险", "RACI、结对评审和运维演练"],
        ["数据", "简历、录音和评价的隐私风险", "最小采集、权限、脱敏、删除和审计"],
    ], [1700, 3760, 3900])
    doc.add_heading("预期收尾与归档要求", level=1)
    add_management_table(doc, ["收尾事项", "责任人", "完成判据"], [
        ["范围与验收确认", "陈昊/赵宁", "Must 交付物与验收证据齐备"],
        ["质量和风险关闭审查", "马琳/韩磊", "严重缺陷为 0，残余风险获批准"],
        ["部署、恢复与知识移交", "韩磊/林泽宇", "运行手册、恢复演练和培训记录齐备"],
        ["合同、预算和资源释放", "陈昊", "采购关闭、预算核对和资源释放获批准"],
        ["经验教训与资料归档", "全体虚构项目团队", "决策、变更、测试、风险和复盘资料可追溯"],
    ], [3900, 2100, 3360])
    doc.add_heading("总结结论", level=1)
    add_callout(doc, "结论", "本套项目管理方案围绕 20 周总计划、13 人虚构团队、8,440 工时和 %s 预算建立了范围、进度、成本、质量、风险与团队管理基线。项目是否成功只能在未来按这些基线和验收证据评价，本报告不预设执行结果。" % money(PROJECT["bac"]))
    return doc


BUILDERS = {
    "01_项目建议书.docx": build_proposal,
    "02_可行性分析报告.docx": build_feasibility,
    "03_项目计划书.docx": build_plan,
    "04_项目估算报告.docx": build_estimate,
    "05_项目进度和成本控制报告.docx": build_control,
    "06_项目质量管理报告.docx": build_quality,
    "07_项目风险管理报告.docx": build_risk,
    "08_项目团队建设和管理报告.docx": build_team,
    "09_项目总结报告.docx": build_summary,
}


def generate_all(output_dir=OUTPUT_DIR):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for filename, _ in REPORTS:
        doc = BUILDERS[filename]()
        path = output_dir / filename
        doc.save(path)
        paths.append(path)
    return paths


if __name__ == "__main__":
    for generated in generate_all():
        print(generated)
