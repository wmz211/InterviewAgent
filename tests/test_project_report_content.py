import tempfile
import unittest
from pathlib import Path

from scripts.audit_project_management_reports import audit_directory
from scripts.generate_project_management_reports import generate_all
from docx import Document


class ProjectReportContentTests(unittest.TestCase):
    def test_all_nine_reports_pass_content_and_structure_audit(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "reports"
            paths = generate_all(output)
            self.assertEqual(len(paths), 9)
            errors = audit_directory(output)
            self.assertEqual(errors, [])
            for path in paths:
                doc = Document(path)
                text = "\n".join(
                    [p.text for p in doc.paragraphs]
                    + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
                )
                self.assertIn("作业编制小组", text)
                self.assertIn("陈昊", text)
                self.assertIn("林泽宇", text)
                for forbidden in (
                    "最终模拟实际",
                    "第 14 周状态日",
                    "结项状态",
                    "实际成本",
                    "较基线延期",
                    "三人全职",
                    "王明哲（项目经理",
                    "江悦铭（后端",
                    "周士斌（前端",
                ):
                    self.assertNotIn(forbidden, text)


if __name__ == "__main__":
    unittest.main()
